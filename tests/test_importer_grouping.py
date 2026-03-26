import json
import tempfile
from pathlib import Path
import unittest

from openpyxl import Workbook

from core.material_manager import MaterialManager


class TestImporterGrouping(unittest.TestCase):
    def _load_default_config(self) -> dict:
        repo_root = Path(__file__).resolve().parents[1]
        with (repo_root / "config" / "default_config.json").open("r", encoding="utf-8") as f:
            return json.load(f)

    def _write_test_workbook(self, path: Path) -> None:
        wb = Workbook()

        # Header names intentionally include whitespace/case variations to
        # ensure Milestone 2 importer matching is robust.
        header = [
            "LOT ID ",
            "DATE",
            "Wafer ID",
            "KW.H",
            "Ar ",  # variation
            "O2",
            "IncidentAngle",  # variation (no space)
            "Linear-Offset",  # variation (dash punctuation)
            "Rotations",
            "rpm(rotation velocity)",
            "Power(W)",
            "Thickness",
            "RS",
            "RSU",
            "Target ID",
        ]

        # Sheet: Ta
        ws_ta = wb.active
        ws_ta.title = "Ta"
        for c, name in enumerate(header, start=1):
            ws_ta.cell(row=1, column=c, value=name)

        # Valid Ta.1 rows (include duplicate lifetime within target instance).
        ws_ta.append(["L1", "2024-01-01", "W1", 0.1, 10, 1, 20.0, 0.5, 100, 60.0, 1000.0, 1000.0, 5.0, 0.2, "Ta.1"])
        ws_ta.append(["L2", "2024-01-02", "W2", 0.1, 10, 1, 21.0, 0.6, 101, 60.0, 1000.0, 1001.0, 5.2, 0.25, "Ta.1"])
        ws_ta.append(["L3", "2024-01-03", "W3", 0.2, 10, 1, 22.0, 0.7, 102, 60.0, 1000.0, 1002.0, 5.1, 0.18, "Ta.1"])

        # Valid Ta.2 row, but with negative RS (suspicious negative).
        ws_ta.append(["L4", "2024-01-04", "W4", 0.3, 10, 1, 23.0, 0.8, 103, 60.0, 1000.0, 1003.0, -1.0, 0.22, "Ta.2"])

        # Blank Target ID row (should trigger target_id warning and be excluded from grouping).
        ws_ta.append(["L5", "2024-01-05", "W5", 0.4, 10, 1, 24.0, 0.9, 104, 60.0, 1000.0, 1004.0, 5.4, 0.21, ""])

        # Sheet: Ru
        ws_ru = wb.create_sheet("Ru")
        for c, name in enumerate(header, start=1):
            ws_ru.cell(row=1, column=c, value=name)

        ws_ru.append(["L6", "2024-01-01", "W6", 0.1, 8, 2, 10.0, 0.2, 50, 60.0, 1000.0, 900.0, 3.3, 0.3, "Ru.1"])
        ws_ru.append(["L7", "2024-01-02", "W7", 0.2, 8, 2, 11.0, 0.25, 52, 60.0, 1000.0, 910.0, 3.1, 0.28, "Ru.1"])

        wb.save(str(path))

    def test_workbook_import_and_target_grouping(self) -> None:
        cfg = self._load_default_config()
        column_map = cfg["column_mapping"]

        with tempfile.TemporaryDirectory() as td:
            tmp_path = Path(td) / "test_workbook.xlsx"
            self._write_test_workbook(tmp_path)

            manager = MaterialManager()
            materials = manager.load_workbook(
                str(tmp_path),
                column_map=column_map,
                config=cfg,
            )

            self.assertEqual(set(materials), {"Ta", "Ru"})

            ta_ds = manager.get_material("Ta")
            self.assertIn("Ta.1", ta_ds.target_instances)
            self.assertIn("Ta.2", ta_ds.target_instances)

            ta1_records = ta_ds.target_instances["Ta.1"].records
            ta2_records = ta_ds.target_instances["Ta.2"].records

            # Ta.1: 3 rows (including the duplicate lifetime row).
            self.assertEqual(len(ta1_records), 3)
            # Ta.2: 1 row.
            self.assertEqual(len(ta2_records), 1)

            # Validate that key canonical columns exist.
            self.assertIn("lifetime", ta_ds.all_records.columns)
            self.assertIn("target_id", ta_ds.all_records.columns)
            self.assertIn("rpm", ta_ds.all_records.columns)
            self.assertIn("power", ta_ds.all_records.columns)

            # Ensure validation warnings were produced.
            warnings_joined = "\n".join(ta_ds.warnings)
            self.assertIn("duplicate_lifetime_within_target_instance", warnings_joined)
            self.assertIn("suspicious_negative_values", warnings_joined)
            self.assertIn("target_id_missing_or_malformed", warnings_joined)


if __name__ == "__main__":
    unittest.main()

